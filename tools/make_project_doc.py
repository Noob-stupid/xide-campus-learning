# -*- coding: utf-8 -*-
"""把 docs/项目详细介绍.md 转成适合提交的 Word 文档。"""
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, 'docs', '项目详细介绍.md')
OUT = sys.argv[2] if len(sys.argv) > 2 else os.path.join(ROOT, 'docs', '习得-项目详细介绍.docx')

FONT = '微软雅黑'
INK = RGBColor(0x18, 0x19, 0x1C)
BLUE = RGBColor(0x00, 0x84, 0xB0)
MUTED = RGBColor(0x61, 0x66, 0x6D)


def set_font(run, size=10.5, bold=False, color=INK, mono=False):
    run.font.name = 'Consolas' if mono else FONT
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_para(doc, text, size=10.5, bold=False, color=INK, mono=False, align=None, space_after=4):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.35
    run = p.add_run(text)
    set_font(run, size=size, bold=bold, color=color, mono=mono)
    return p


def add_table(doc, rows):
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncol)
    table.style = 'Light Grid Accent 1'
    for i, row in enumerate(rows):
        for j in range(ncol):
            cell = table.cell(i, j)
            cell.text = ''
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            text = row[j] if j < len(row) else ''
            run = p.add_run(text)
            set_font(run, size=9, bold=(i == 0), color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def clean(s):
    s = s.replace('`', '')
    s = re.sub(r'\*\*(.+?)\*\*', r'\1', s)
    return s.strip()


def main():
    with open(SRC, encoding='utf-8') as fh:
        lines = fh.read().splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.4)
    sec.right_margin = Cm(2.4)

    normal = doc.styles['Normal']
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    normal.font.size = Pt(10.5)

    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hr = header.add_run('习得 · 校园学习空间 AI 创新应用　|　项目详细介绍')
    set_font(hr, size=8.5, color=MUTED)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run('习得 · 从落座到学会　|　离座，不放学。')
    set_font(fr, size=8, color=MUTED)

    i = 0
    table_buf = []
    code_buf = []
    in_code = False
    while i < len(lines):
        line = lines[i]
        if line.strip().startswith('```'):
            if in_code:
                if code_buf:
                    for cl in code_buf:
                        add_para(doc, cl, size=8.5, mono=True, color=MUTED, space_after=1)
                    doc.add_paragraph().paragraph_format.space_after = Pt(2)
                code_buf = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue
        if line.strip().startswith('|') and line.strip().endswith('|'):
            cells = [clean(c) for c in line.strip().strip('|').split('|')]
            if all(re.fullmatch(r':?-{2,}:?', c.replace(' ', '')) for c in cells):
                i += 1
                continue
            table_buf.append(cells)
            i += 1
            continue
        if table_buf:
            add_table(doc, table_buf)
            table_buf = []
        s = line.rstrip()
        if not s.strip():
            i += 1
            continue
        if s.startswith('# '):
            add_para(doc, clean(s[2:]), size=22, bold=True, color=BLUE, space_after=10)
        elif s.startswith('## '):
            add_para(doc, clean(s[3:]), size=15, bold=True, color=BLUE, space_after=6)
        elif s.startswith('### '):
            add_para(doc, clean(s[4:]), size=12, bold=True, color=INK, space_after=4)
        elif s.strip() == '---':
            doc.add_paragraph()
        elif s.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(clean(s[2:]))
            set_font(run, size=10)
        elif re.match(r'^\d+\.\s+', s):
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(clean(re.sub(r'^\d+\.\s+', '', s)))
            set_font(run, size=10)
        elif s.startswith('> '):
            add_para(doc, clean(s[2:]), size=10, color=MUTED)
        else:
            add_para(doc, clean(s), size=10.5)
        i += 1
    if table_buf:
        add_table(doc, table_buf)
    doc.save(OUT)
    print('saved', OUT, os.path.getsize(OUT), 'bytes')


if __name__ == '__main__':
    main()
